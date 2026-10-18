"""Exercise source readers with real file fixtures and isolated external connections."""

from contextlib import redirect_stderr, redirect_stdout
from email.message import EmailMessage
from io import StringIO
import mailbox
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch

from docx import Document as WordDocument
from openpyxl import Workbook
import requests

import main
from src.sources import load_documents


class SourceReaderTests(unittest.TestCase):
    """Check extraction and provenance through the public source-loading interface."""

    def setUp(self):
        """Create a fresh source folder for each test and remove it afterward."""
        temporary = TemporaryDirectory(prefix="rag-sources-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def write_text(self, name, content):
        """Write a UTF-8 fixture and return its path."""
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return path

    def response(self, content, content_type, url="https://example.test/source"):
        """Provide HTTP response bytes without reaching a real remote website."""
        response = MagicMock()
        response.__enter__.return_value = response
        response.iter_content.return_value = [content]
        response.headers = {"Content-Type": content_type}
        response.url = url
        return response

    def test_text_and_code_are_read_without_execution(self):
        """Plain text, Markdown, and code retain their original text and filename."""
        for name in ("guide.txt", "guide.md", "example.py"):
            with self.subTest(name=name):
                path = self.write_text(name, "raise RuntimeError('Read this; do not execute it')")
                documents = load_documents(path)
                self.assertIn("raise RuntimeError", documents[0].page_content)
                self.assertEqual(documents[0].metadata["source"], str(path.resolve()))

    def test_csv_and_tsv_keep_headers_with_each_row(self):
        """Table records repeat column labels and preserve row numbers."""
        for extension, delimiter in (("csv", ","), ("tsv", "\t")):
            with self.subTest(extension=extension):
                path = self.write_text(
                    f"policy.{extension}",
                    f"product{delimiter}warranty\nCedar{delimiter}37 days\n\n",
                )
                documents = load_documents(path)
                self.assertEqual(len(documents), 1)
                self.assertIn("product: Cedar", documents[0].page_content)
                self.assertIn("warranty: 37 days", documents[0].page_content)
                self.assertEqual(documents[0].metadata["row"], 2)

    def test_word_preserves_paragraph_and_table_order(self):
        """Word tables stay between the paragraphs surrounding them."""
        path = self.root / "manual.docx"
        word = WordDocument()
        word.add_paragraph("Before the table")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Warranty: 37 days"
        word.add_paragraph("After the table")
        word.save(path)
        content = load_documents(path)[0].page_content
        self.assertLess(content.index("Before"), content.index("Warranty"))
        self.assertLess(content.index("Warranty"), content.index("After"))

    def test_excel_reads_all_sheets_with_sheet_metadata(self):
        """Each Excel sheet produces rows with its own name and column labels."""
        path = self.root / "policies.xlsx"
        workbook = Workbook()
        for sheet in (workbook.active, workbook.create_sheet("Returns")):
            sheet.append(["Product", "Days"])
            sheet.append(["Cedar", 37])
        workbook.save(path)
        workbook.close()
        documents = load_documents(path)
        self.assertEqual(len(documents), 2)
        self.assertEqual({doc.metadata["sheet"] for doc in documents}, {"Sheet", "Returns"})
        self.assertTrue(all("Days: 37" in doc.page_content for doc in documents))

    def test_json_and_json_lines_preserve_records(self):
        """JSON arrays and JSONL become separate searchable records."""
        fixtures = {
            "items.json": '[{"name": "Cedar"}, {"name": "Maple"}]',
            "items.jsonl": '{"name": "Cedar"}\n{"name": "Maple"}\n',
        }
        for name, content in fixtures.items():
            with self.subTest(name=name):
                documents = load_documents(self.write_text(name, content))
                self.assertEqual(len(documents), 2)
                self.assertIn("Maple", documents[1].page_content)
                self.assertEqual(documents[1].metadata["record"], 2)

    def test_html_keeps_visible_text_and_removes_scripts(self):
        """HTML contributes readable content and a title, excluding executable scripts."""
        path = self.write_text(
            "page.html",
            "<title>Manual</title><script>hidden script</script><p>Warranty: 37 days</p>",
        )
        document = load_documents(path)[0]
        self.assertIn("Warranty: 37 days", document.page_content)
        self.assertNotIn("hidden script", document.page_content)
        self.assertEqual(document.metadata["title"], "Manual")

    def test_email_reads_body_without_attachments(self):
        """EML files include message headers and body while excluding attachments."""
        message = EmailMessage()
        message["Subject"] = "Warranty policy"
        message["From"] = "support@example.test"
        message.set_content("The warranty lasts 37 days.")
        message.add_attachment("Attachment should not be indexed", filename="attachment.txt")
        path = self.root / "message.eml"
        path.write_bytes(message.as_bytes())
        content = load_documents(path)[0].page_content
        self.assertIn("Subject: Warranty policy", content)
        self.assertIn("37 days", content)
        self.assertNotIn("Attachment should not", content)

    def test_html_email_falls_back_to_visible_text(self):
        """HTML-only email bodies are converted to readable text."""
        message = EmailMessage()
        message.set_content("<p>Warranty: 37 days</p>", subtype="html")
        path = self.root / "message.eml"
        path.write_bytes(message.as_bytes())
        content = load_documents(path)[0].page_content
        self.assertIn("Warranty: 37 days", content)
        self.assertNotIn("<p>", content)

    def test_mbox_reads_multiple_exported_messages(self):
        """Mailbox exports preserve each message as a separate document."""
        path = self.root / "archive.mbox"
        archive = mailbox.mbox(path)
        try:
            for subject in ("Warranty", "Returns"):
                message = EmailMessage()
                message["Subject"] = subject
                message.set_content(f"{subject} window: 37 days")
                archive.add(message)
            archive.flush()
        finally:
            archive.close()
        documents = load_documents(path)
        self.assertEqual(len(documents), 2)
        self.assertIn("Returns", documents[1].page_content)
        self.assertEqual(documents[1].metadata["message_index"], 2)

    def test_mixed_directory_recurses_and_skips_generated_files(self):
        """Mixed folders include supported sources and exclude hidden/generated content."""
        self.write_text("nested/guide.MD", "Manual")
        self.write_text("data.csv", "name,days\nCedar,37")
        self.write_text(".hidden/secret.txt", "hidden")
        self.write_text("node_modules/library.js", "generated")
        self.write_text("image.png", "unsupported")
        documents = load_documents(self.root)
        self.assertEqual(len(documents), 2)
        self.assertEqual({Path(doc.metadata["source"]).name for doc in documents}, {"guide.MD", "data.csv"})

    def test_invalid_and_unsupported_sources_are_actionable(self):
        """Errors identify missing, unsupported, or malformed files."""
        with self.assertRaisesRegex(ValueError, "does not exist"):
            load_documents(self.root / "missing.txt")
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            load_documents(self.write_text("old.doc", "legacy format"))
        with self.assertRaisesRegex(ValueError, "broken.json"):
            load_documents(self.write_text("broken.json", "invalid JSON"))
        with self.assertRaisesRegex(ValueError, "HTTP"):
            load_documents("ftp://example.test/manual.txt")

    def test_url_reads_html_and_retains_the_original_url(self):
        """A public page is searchable by its visible text with URL provenance."""
        url = "https://example.test/manual"
        response = self.response(b"<title>Manual</title><p>Warranty: 37 days</p>", "text/html; charset=utf-8", url)
        with patch("src.sources.requests.get", return_value=response):
            document = load_documents(url)[0]
        self.assertIn("37 days", document.page_content)
        self.assertEqual(document.metadata["source"], url)

    def test_url_reads_downloadable_word_files(self):
        """Downloaded documents use their HTTP format even when the URL lacks a suffix."""
        path = self.root / "manual.docx"
        word = WordDocument()
        word.add_paragraph("Warranty: 37 days")
        word.save(path)
        response = self.response(
            path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        with patch("src.sources.requests.get", return_value=response):
            document = load_documents("https://example.test/download")[0]
        self.assertIn("37 days", document.page_content)
        self.assertEqual(document.metadata["source"], "https://example.test/download")

    def test_url_reports_http_errors_and_download_limits(self):
        """Network failures and excessive responses become readable source errors."""
        response = self.response(b"too large", "text/plain")
        with patch("src.sources.requests.get", return_value=response), patch("src.sources.MAX_URL_BYTES", 3):
            with self.assertRaisesRegex(ValueError, "download limit"):
                load_documents("https://example.test/large")
        with patch("src.sources.requests.get", side_effect=requests.Timeout("timed out")):
            with self.assertRaisesRegex(ValueError, "Could not load URL"):
                load_documents("https://example.test/unavailable")

    def test_sqlite_select_extracts_rows_without_changing_data(self):
        """SQLite queries create labeled row documents and leave the database intact."""
        path = self.root / "catalog.sqlite"
        with sqlite3.connect(path) as connection:
            connection.execute("CREATE TABLE products (name TEXT, days INTEGER)")
            connection.execute("INSERT INTO products VALUES ('Cedar', 37)")
        url = f"sqlite:///{path}"
        documents = load_documents(url, sql="SELECT name, days FROM products")
        self.assertIn("name: Cedar", documents[0].page_content)
        self.assertIn("days: 37", documents[0].page_content)
        self.assertEqual(documents[0].metadata["source_type"], "database")
        with self.assertRaisesRegex(ValueError, "SELECT query"):
            load_documents(url, sql="DELETE FROM products")
        with self.assertRaisesRegex(ValueError, "query failed"):
            load_documents(url, sql="SELECT * FROM products; DELETE FROM products")
        with sqlite3.connect(path) as connection:
            self.assertEqual(connection.execute("SELECT count(*) FROM products").fetchone()[0], 1)

    def test_missing_sqlite_database_is_not_created(self):
        """A mistyped SQLite location must not silently create an empty database."""
        path = self.root / "missing.sqlite"
        with self.assertRaisesRegex(ValueError, "existing database file"):
            load_documents(f"sqlite:///{path}", sql="SELECT 1")
        self.assertFalse(path.exists())

    def test_server_databases_enable_read_only_and_omit_credentials(self):
        """Server adapters select a read-only transaction before executing the query."""
        commands = {
            "postgresql": "SET TRANSACTION READ ONLY",
            "mysql": "START TRANSACTION READ ONLY",
        }
        for dialect, readonly in commands.items():
            with self.subTest(dialect=dialect), patch("src.sources.create_engine") as factory:
                connection = factory.return_value.connect.return_value.__enter__.return_value
                result = connection.execute.return_value
                result.keys.return_value = ["name"]
                result.__iter__.return_value = iter([("Cedar",)])
                documents = load_documents(
                    f"{dialect}://reader:secret@localhost/catalog", sql="SELECT name FROM products"
                )
                connection.exec_driver_sql.assert_called_once_with(readonly)
                self.assertIn("Cedar", documents[0].page_content)
                self.assertNotIn("secret", documents[0].metadata["source"])
                self.assertNotIn("reader", documents[0].metadata["source"])


class SourceCliTests(unittest.TestCase):
    """Check source selection and option validation at the command-line interface."""

    def test_cli_selects_local_url_and_database_sources(self):
        """Each source flag reaches the shared ingestion function with correct arguments."""
        cases = [
            (["--path", "manual.docx"], Path("manual.docx"), None),
            (["--url", "https://example.test/manual"], "https://example.test/manual", None),
            (["--database-url", "sqlite:///catalog.db", "--sql", "SELECT * FROM products"], "sqlite:///catalog.db", "SELECT * FROM products"),
        ]
        for options, source, sql in cases:
            with self.subTest(options=options), patch("sys.argv", ["main.py", "ingest", *options]):
                with patch("src.ingest.ingest") as ingest:
                    main.main()
                ingest.assert_called_once_with(source, sql=sql)

    def test_cli_rejects_incomplete_or_conflicting_source_options(self):
        """SQL needs a database URL, and a command accepts one source selection."""
        for options in (
            ["--sql", "SELECT 1"],
            ["--database-url", "sqlite:///catalog.db"],
            ["--path", "manual.txt", "--url", "https://example.test/manual"],
        ):
            with self.subTest(options=options), patch("sys.argv", ["main.py", "ingest", *options]):
                with redirect_stderr(StringIO()), redirect_stdout(StringIO()):
                    with self.assertRaises(SystemExit) as error:
                        main.main()
                self.assertEqual(error.exception.code, 2)


if __name__ == "__main__":
    unittest.main()
