"""Read different sources through one interface that returns LangChain documents.

Format-specific parsing lives here. Ingestion and retrieval work with the same
Document objects regardless of whether their text came from a file, URL, or SQL.
"""

import csv
from email import policy
from email.parser import BytesParser
from itertools import zip_longest
import json
import mailbox
import os
from pathlib import Path
import re
import sqlite3
from tempfile import TemporaryDirectory
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.text.paragraph import Paragraph
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from openpyxl import load_workbook
import requests
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL, make_url
from sqlalchemy.exc import SQLAlchemyError


# Plain text and source code are read as UTF-8 without executing their contents.
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".xml", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".java", ".c", ".cpp", ".h", ".hpp", ".go", ".rs", ".rb", ".php",
    ".sh", ".sql", ".css", ".scss", ".swift", ".kt", ".ipynb",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {
    ".pdf", ".docx", ".csv", ".tsv", ".xlsx", ".json", ".jsonl",
    ".html", ".htm", ".eml", ".mbox",
}
SKIP_DIRECTORIES = {"node_modules", "__pycache__", "vector_store"}
MAX_URL_BYTES = 20 * 1024 * 1024


def _document(content: str, source: str, source_type: str, **metadata) -> Document:
    """Attach consistent source metadata to a piece of extracted text."""
    return Document(
        page_content=content,
        metadata={"source": source, "source_type": source_type, **metadata},
    )


def _html_text(content: str | bytes) -> tuple[str, str]:
    """Extract visible HTML text and its title, omitting scripts and styles."""
    soup = BeautifulSoup(content, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    return soup.get_text("\n", strip=True), title


def _table_documents(rows, source: str, source_type: str, **metadata) -> list[Document]:
    """Turn table rows into documents using the first nonempty row as headers.

    Column labels are repeated for every record so a retrieved row retains
    its meaning even when other rows are not included in the prompt.
    """
    rows = iter(rows)
    headers = None
    documents = []
    for row_number, row in enumerate(rows, start=1):
        values = ["" if value is None else str(value) for value in row]
        if not any(value.strip() for value in values):
            continue
        if headers is None:
            headers = values
            continue
        lines = [
            f"{header or f'Column {index}'}: {value}"
            for index, (header, value) in enumerate(
                zip_longest(headers, values, fillvalue=""), start=1
            )
        ]
        if "sheet" in metadata:
            lines.insert(0, f"Sheet: {metadata['sheet']}")
        documents.append(_document(
            "\n".join(lines), source, source_type, row=row_number, **metadata
        ))
    return documents


def _email_document(message, source: str, **metadata) -> Document:
    """Read email headers and the preferred text body, excluding attachments."""
    body = message.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body is not None else ""
    if body is not None and body.get_content_subtype() == "html":
        content, _ = _html_text(content)
    headers = "\n".join(
        f"{name}: {message[name]}"
        for name in ("Subject", "From", "To", "Date") if message[name]
    )
    return _document(
        f"{headers}\n\n{content}".strip(), source, "email", **metadata
    )


def _load_file(path: Path) -> list[Document]:
    """Dispatch a supported file to its reader while preserving source details."""
    source = str(path)
    extension = path.suffix.lower()
    if extension == ".pdf":
        documents = PyPDFLoader(source).load()
        for document in documents:
            document.metadata["source_type"] = "pdf"
        return documents
    if extension in TEXT_EXTENSIONS:
        return [_document(path.read_text(encoding="utf-8-sig"), source, "text")]
    if extension in {".html", ".htm"}:
        content, title = _html_text(path.read_bytes())
        return [_document(content, source, "html", title=title)]
    if extension == ".docx":
        word = WordDocument(path)
        parts = []
        # Preserve the order of paragraphs and tables from the Word document.
        for block in word.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            else:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        return [_document("\n".join(parts), source, "docx")]
    if extension in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = csv.reader(handle, delimiter="\t" if extension == ".tsv" else ",")
            return _table_documents(rows, source, extension[1:])
    if extension == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            documents = []
            for sheet in workbook.worksheets:
                documents.extend(_table_documents(
                    sheet.iter_rows(values_only=True), source, "xlsx", sheet=sheet.title
                ))
            return documents
        finally:
            workbook.close()
    if extension in {".json", ".jsonl"}:
        content = path.read_text(encoding="utf-8-sig")
        if extension == ".jsonl":
            records = [json.loads(line) for line in content.splitlines() if line.strip()]
        else:
            data = json.loads(content)
            records = data if isinstance(data, list) else [data]
        return [
            _document(json.dumps(record, ensure_ascii=False, indent=2), source, "json", record=index)
            for index, record in enumerate(records, start=1)
        ]
    if extension == ".eml":
        with path.open("rb") as handle:
            message = BytesParser(policy=policy.default).parse(handle)
        return [_email_document(message, source)]
    if extension == ".mbox":
        archive = mailbox.mbox(
            path, create=False,
            factory=lambda handle: BytesParser(policy=policy.default).parse(handle),
        )
        try:
            return [
                _email_document(message, source, message_index=index)
                for index, message in enumerate(archive, start=1)
            ]
        finally:
            archive.close()
    raise ValueError(f"Unsupported file type '{extension or '(no extension)'}': {path}")


def _source_files(folder: Path):
    """Find supported files recursively, skipping hidden and generated folders."""
    for root, directories, names in os.walk(folder):
        directories[:] = sorted(
            name for name in directories
            if not name.startswith(".") and name not in SKIP_DIRECTORIES
        )
        for name in sorted(names):
            path = Path(root) / name
            if not name.startswith(".") and not path.is_symlink() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                yield path


def _load_path(path: Path) -> list[Document]:
    """Read one local file or all supported files in a directory."""
    path = path.expanduser().resolve()
    if not path.exists():
        raise ValueError(f"Source path does not exist: {path}")
    files = list(_source_files(path)) if path.is_dir() else [path]
    documents = []
    for file in files:
        try:
            documents.extend(_load_file(file))
        except Exception as error:
            # Include the failing file when a mixed folder contains a corrupt source.
            raise ValueError(f"Could not read {file}: {error}") from error
    return documents


def _load_url(url: str) -> list[Document]:
    """Download one public page or supported document; retain the URL as source.

    This reads server-provided content without running JavaScript, logging in,
    or following links to other pages. Downloads are limited to 20 MiB.
    """
    try:
        with requests.get(
            url, timeout=(10, 30), stream=True,
            headers={"User-Agent": "local-rag/1.0"},
        ) as response:
            response.raise_for_status()
            chunks = []
            size = 0
            for chunk in response.iter_content(chunk_size=65536):
                size += len(chunk)
                if size > MAX_URL_BYTES:
                    raise ValueError("URL content exceeds the 20 MiB download limit.")
                chunks.append(chunk)
            content = b"".join(chunks)
            content_type = response.headers.get("Content-Type", "").split(";")[0].lower().strip()
            extension = Path(urlparse(response.url).path).suffix.lower()
            types = {
                "text/html": ".html", "application/xhtml+xml": ".html",
                "text/plain": ".txt", "text/csv": ".csv",
                "text/tab-separated-values": ".tsv", "application/json": ".json",
                "application/pdf": ".pdf",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            }
            # HTTP's declared format takes precedence over a misleading URL suffix.
            extension = types.get(content_type, extension)
            if extension not in SUPPORTED_EXTENSIONS:
                raise ValueError(f"Unsupported URL content type: {content_type or 'unknown'}")
            with TemporaryDirectory(prefix="rag-download-") as temporary:
                path = Path(temporary) / f"source{extension}"
                path.write_bytes(content)
                documents = _load_file(path)
            for document in documents:
                document.metadata["source"] = url
            return documents
    except requests.RequestException as error:
        raise ValueError(f"Could not load URL {url}: {error}") from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"Could not parse document from {url}: {error}") from error


def _load_database(database_url: str, sql: str) -> list[Document]:
    """Read a SELECT result from SQLite, PostgreSQL, or MySQL as row documents.

    Database-level read-only settings prevent writes during ingestion. Only
    selected rows are indexed; database credentials are omitted from metadata.
    """
    if not re.match(r"SELECT\b", sql.strip(), flags=re.IGNORECASE):
        raise ValueError("Database ingestion requires a SELECT query.")
    try:
        url = make_url(database_url)
    except SQLAlchemyError as error:
        raise ValueError("Invalid database URL.") from error
    dialect = url.get_backend_name()
    if dialect not in {"sqlite", "postgresql", "mysql"}:
        raise ValueError("Supported databases are SQLite, PostgreSQL, and MySQL.")
    if dialect == "sqlite":
        path = Path(url.database or "").expanduser().resolve()
        if not path.is_file():
            raise ValueError("SQLite ingestion requires an existing database file.")
        # SQLite's read-only URI also prevents accidentally creating a new database.
        engine = create_engine(
            "sqlite://",
            creator=lambda: sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True),
        )
    else:
        # Select the drivers installed by requirements.txt for ordinary URLs too.
        driver = "postgresql+psycopg" if dialect == "postgresql" else "mysql+pymysql"
        url = url.set(drivername=driver)
        engine = create_engine(url)
    source = URL.create(
        drivername=url.drivername, host=url.host, port=url.port, database=url.database,
    ).render_as_string()
    try:
        with engine.connect() as connection:
            if dialect == "sqlite":
                connection.exec_driver_sql("PRAGMA query_only = ON")
            elif dialect == "postgresql":
                connection.exec_driver_sql("SET TRANSACTION READ ONLY")
            else:
                connection.exec_driver_sql("START TRANSACTION READ ONLY")
            result = connection.execute(text(sql))
            rows = [tuple(result.keys()), *result]
            return _table_documents(rows, source, "database")
    except SQLAlchemyError as error:
        raise ValueError(f"Read-only database query failed: {error}") from error
    finally:
        engine.dispose()


def load_documents(source: str | Path, *, sql: str | None = None) -> list[Document]:
    """Read a local file/folder, an HTTP(S) URL, or an explicit database SELECT.

    Args:
        source: File/folder path, web URL, or SQLAlchemy database URL.
        sql: SELECT statement; providing it selects database ingestion.

    Returns:
        Documents with text and source metadata, ready for the shared splitter.

    Raises:
        ValueError: The source is invalid, unreadable, or unsupported.
    """
    if sql is not None:
        return _load_database(str(source), sql)
    if isinstance(source, str) and urlparse(source).scheme.lower() in {"http", "https"}:
        return _load_url(source)
    if isinstance(source, str) and "://" in source:
        raise ValueError("Use an HTTP(S) URL, or provide --database-url and --sql for a database.")
    return _load_path(Path(source))
